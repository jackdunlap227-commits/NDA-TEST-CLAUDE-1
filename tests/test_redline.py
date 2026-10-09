import copy
import io
import sys
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "samples"))
from make_sample_nda import build  # noqa: E402

from nda_agent.redline import Edit, RedlineDocument  # noqa: E402


@pytest.fixture(scope="module")
def sample_bytes(tmp_path_factory) -> bytes:
    path = tmp_path_factory.mktemp("nda") / "nda.docx"
    build(path)
    return path.read_bytes()


def _view(data: bytes, mode: str) -> list[str]:
    """Paragraph texts with all changes accepted (mode='accept') or rejected ('reject')."""
    body = copy.deepcopy(docx.Document(io.BytesIO(data)).element.body)
    drop, keep = (qn("w:del"), qn("w:ins")) if mode == "accept" else (qn("w:ins"), qn("w:del"))
    for el in list(body.iter(drop)):
        if el.getparent().tag != qn("w:rPr"):
            el.getparent().remove(el)
    out = []
    for p in body.iter(qn("w:p")):
        ppr_mark = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:ins')}")
        if mode == "reject" and ppr_mark is not None:
            continue
        text = "".join(t.text or "" for t in p.iter(qn("w:t"), qn("w:delText")))
        out.append(text)
    return out


def test_replace_across_runs_with_smart_quotes(sample_bytes):
    doc = RedlineDocument(sample_bytes)
    original = doc.paragraph_text("p4")
    # The model quotes with straight quotes; the doc has curly ones. Target spans two runs.
    edits = [
        Edit("p4", "replace", "officers, directors, employees and attorneys",
             "and its affiliates' officers, directors, employees, agents, attorneys, consultants, "
             "potential financing sources and advisors", "R3 Representatives"),
        Edit("p4", "insert", "provided by you", " or on your behalf", "R4 Representatives"),
    ]
    result = doc.apply(edits)
    assert not result.skipped, result.skipped
    out = doc.save()

    accepted = _view(out, "accept")[4]
    assert "potential financing sources and advisors to whom" in accepted
    assert accepted.endswith("provided by you or on your behalf.")
    assert _view(out, "reject")[4] == original


def test_heading_run_split(sample_bytes):
    # "Representatives." heading is a separate bold run from the body text.
    doc = RedlineDocument(sample_bytes)
    result = doc.apply([Edit("p4", "replace", "Representatives. “Representatives”", "Reps. “Reps”")])
    assert not result.skipped
    assert _view(doc.save(), "accept")[4].startswith("3. Reps. “Reps” means")


def test_pure_deletion_and_new_paragraph(sample_bytes):
    doc = RedlineDocument(sample_bytes)
    n_before = len(_view(sample_bytes, "accept"))
    result = doc.apply([
        Edit("p6", "replace", "immediate prior written ", "", "R12"),
        Edit("p3", "insert_paragraph_after", "", "(iii) is independently developed by Recipient.", "R7"),
    ])
    assert not result.skipped
    out = doc.save()
    accepted = _view(out, "accept")
    assert len(accepted) == n_before + 1
    assert "provide the Company with notice" in accepted[7]
    assert accepted[4] == "(iii) is independently developed by Recipient."
    assert _view(out, "reject") == _view(sample_bytes, "accept")


def test_unmatched_edit_is_skipped_and_commented(sample_bytes):
    doc = RedlineDocument(sample_bytes)
    result = doc.apply([Edit("p5", "replace", "this text does not exist", "x", "R10: note")])
    assert len(result.applied) == 0 and len(result.skipped) == 1
    out = doc.save()
    comments = docx.Document(io.BytesIO(out)).comments
    assert any("Not auto-applied" in c.text for c in comments)
    assert _view(out, "accept") == _view(sample_bytes, "accept")


def test_comments_and_track_changes_setting(sample_bytes):
    doc = RedlineDocument(sample_bytes, author="Unifi Legal")
    doc.apply([Edit("p11", "replace", "survive indefinitely",
                    "terminate on the first anniversary of the Effective Date", "R44 Term: one year term")])
    out = doc.save()
    d = docx.Document(io.BytesIO(out))
    assert [c.text for c in d.comments] == ["R44 Term: one year term"]
    assert d.settings.element.find(qn("w:trackRevisions")) is not None
    dels = list(d.element.body.iter(qn("w:del")))
    assert dels and all(el.get(qn("w:author")) == "Unifi Legal" for el in dels)
    assert "".join(t.text for el in dels for t in el.iter(qn("w:delText"))) == "survive indefinitely"


def test_numbered_text_skips_existing_deletions(sample_bytes):
    doc = RedlineDocument(sample_bytes)
    doc.apply([Edit("p11", "replace", "indefinitely", "for one year")])
    reread = RedlineDocument(doc.save())
    assert "indefinitely" not in reread.paragraph_text("p11")
    assert reread.paragraph_text("p11").endswith("survive for one year.")

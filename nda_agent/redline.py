"""Read an NDA .docx and write edits back into it as Word tracked changes.

Deletions become ``<w:del>`` runs (rendered by Word as red strikethrough) and
insertions become ``<w:ins>`` runs (underlined), so the counterparty can
accept or reject each change in Word. Each edit can carry a margin comment
explaining which playbook rule drove it.
"""

from __future__ import annotations

import copy
import datetime as dt
import io
import itertools
import re
from dataclasses import dataclass, field

import docx
from docx.oxml.ns import qn
from lxml import etree

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# Characters that commonly differ between what the model quotes and what the
# document contains (smart quotes, non-breaking spaces, dash variants).
_CHAR_FOLD = {
    "‘": "'", "’": "'", "‚": "'", "‛": "'", "′": "'",
    "“": '"', "”": '"', "„": '"', "‟": '"', "″": '"',
    "‐": "-", "‑": "-", "‒": "-", "–": "-", "—": "-", "−": "-",
    " ": " ", " ": " ", " ": " ", "\t": " ", "\n": " ",
}


@dataclass
class Edit:
    """One change to make in the document.

    kind:
      * ``replace`` - strike ``target_text`` and insert ``new_text`` in its place
        (``new_text`` empty means a pure deletion).
      * ``insert`` - insert ``new_text`` immediately after ``target_text``
        (``target_text`` empty means append at the end of the paragraph).
      * ``insert_paragraph_after`` - add ``new_text`` as a new paragraph after
        the paragraph; ``target_text`` is ignored.
    """

    paragraph_id: str
    kind: str
    target_text: str
    new_text: str
    comment: str = ""


@dataclass
class ApplyResult:
    applied: list[Edit] = field(default_factory=list)
    skipped: list[tuple[Edit, str]] = field(default_factory=list)


class RedlineDocument:
    """A loaded .docx whose body paragraphs are addressable by id (``p0``, ``p1``...)."""

    def __init__(self, data: bytes, author: str = "Unifi Legal (NDA Agent)", initials: str = "UL"):
        self.doc = docx.Document(io.BytesIO(data))
        self.author = author
        self.initials = initials
        self._date = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        self._ids = itertools.count(self._max_revision_id() + 1)
        body = self.doc.element.body
        self.paragraphs: dict[str, etree._Element] = {
            f"p{i}": p for i, p in enumerate(body.iter(qn("w:p")))
        }

    # ------------------------------------------------------------------ reading

    def paragraph_text(self, pid: str) -> str:
        return "".join(text for _, text in _text_runs(self.paragraphs[pid]))

    def numbered_text(self) -> str:
        """The document as ``[pN] text`` lines, which is what the model reviews."""
        lines = []
        for pid in self.paragraphs:
            text = self.paragraph_text(pid)
            if text.strip():
                lines.append(f"[{pid}] {text}")
        return "\n".join(lines)

    # ------------------------------------------------------------------ writing

    def apply(self, edits: list[Edit], add_comments: bool = True) -> ApplyResult:
        result = ApplyResult()
        for edit in edits:
            p = self.paragraphs.get(edit.paragraph_id)
            if p is None:
                result.skipped.append((edit, f"unknown paragraph {edit.paragraph_id}"))
                continue
            try:
                if edit.kind == "insert_paragraph_after":
                    if not edit.new_text.strip():
                        raise _EditError("no text to insert")
                    touched = [self._insert_paragraph_after(p, edit.new_text)]
                elif edit.kind == "replace":
                    touched = self._replace(p, edit.target_text, edit.new_text)
                elif edit.kind == "insert":
                    touched = self._insert(p, edit.target_text, edit.new_text)
                else:
                    raise _EditError(f"unknown edit kind {edit.kind!r}")
            except _EditError as exc:
                result.skipped.append((edit, str(exc)))
                if add_comments and edit.comment:
                    # Still surface the reviewer's point so nothing is silently lost.
                    self._comment_paragraph(
                        p, f"[Not auto-applied: {exc}] Suggested change: "
                        f"{_describe(edit)}\n{edit.comment}"
                    )
                continue
            result.applied.append(edit)
            if add_comments and edit.comment and touched:
                self._add_comment(touched[0], touched[-1], edit.comment)
        return result

    def save(self) -> bytes:
        self._enable_track_changes()
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()

    # ------------------------------------------------------------------ edit ops

    def _replace(self, p, target: str, new_text: str) -> list:
        if not target:
            raise _EditError("replace needs target_text")
        start, end = _locate(p, target)
        runs = _isolate(p, start, end)
        if not runs:
            raise _EditError("target text not found in runs")
        dels = [self._wrap_deleted(r) for r in runs]
        touched = list(dels)
        if new_text:
            ins = self._ins_element(new_text, runs[-1])
            _insert_after(_top_level(dels[-1], p), ins, p)
            touched.append(ins)
        return touched

    def _insert(self, p, anchor: str, new_text: str) -> list:
        if not new_text:
            raise _EditError("insert needs new_text")
        runs = _text_runs(p)
        if not runs:
            raise _EditError("paragraph has no text to anchor to")
        if anchor:
            _, end = _locate(p, anchor)
            before = _isolate(p, 0, end)
            ref = before[-1] if before else runs[-1][0]
        else:
            ref = runs[-1][0]
        ins = self._ins_element(new_text, ref)
        _insert_after(_top_level(ref, p), ins, p)
        return [ins]

    def _insert_paragraph_after(self, p, text: str):
        new_p = etree.Element(qn("w:p"))
        p.addnext(new_p)
        ppr = p.find(qn("w:pPr"))
        if ppr is not None:
            new_ppr = copy.deepcopy(ppr)
            for tag in ("w:sectPr", "w:pPrChange", "w:rPr"):
                for el in new_ppr.findall(qn(tag)):
                    new_ppr.remove(el)
            new_p.append(new_ppr)
        else:
            new_ppr = etree.SubElement(new_p, qn("w:pPr"))
        # Mark the paragraph mark itself as inserted so rejecting removes the paragraph.
        mark_rpr = etree.SubElement(new_ppr, qn("w:rPr"))
        mark_rpr.append(self._revision_el("w:ins"))
        ref_runs = _text_runs(p)
        ins = self._ins_element(text, ref_runs[0][0] if ref_runs else None)
        new_p.append(ins)
        return ins

    # ------------------------------------------------------------------ xml builders

    def _revision_el(self, tag: str):
        el = etree.Element(qn(tag))
        el.set(qn("w:id"), str(next(self._ids)))
        el.set(qn("w:author"), self.author)
        el.set(qn("w:date"), self._date)
        return el

    def _wrap_deleted(self, run):
        wrapper = self._revision_el("w:del")
        run.addprevious(wrapper)
        wrapper.append(run)
        for t in run.findall(qn("w:t")):
            t.tag = qn("w:delText")
        for t in run.findall(qn("w:instrText")):
            t.tag = qn("w:delInstrText")
        return wrapper

    def _ins_element(self, text: str, ref_run):
        wrapper = self._revision_el("w:ins")
        run = etree.SubElement(wrapper, qn("w:r"))
        if ref_run is not None:
            rpr = ref_run.find(qn("w:rPr"))
            if rpr is not None:
                rpr = copy.deepcopy(rpr)
                for tag in ("w:rPrChange", "w:ins", "w:del"):
                    for el in rpr.findall(qn(tag)):
                        rpr.remove(el)
                run.append(rpr)
        t = etree.SubElement(run, qn("w:t"))
        t.text = text
        t.set(XML_SPACE, "preserve")
        return wrapper

    def _add_comment(self, first, last, text: str):
        comment = self.doc.comments.add_comment(text=text, author=self.author, initials=self.initials)
        cid = str(comment.comment_id)
        start = etree.Element(qn("w:commentRangeStart"))
        start.set(qn("w:id"), cid)
        end = etree.Element(qn("w:commentRangeEnd"))
        end.set(qn("w:id"), cid)
        ref_run = etree.Element(qn("w:r"))
        ref = etree.SubElement(ref_run, qn("w:commentReference"))
        ref.set(qn("w:id"), cid)
        first.addprevious(start)
        last.addnext(end)
        end.addnext(ref_run)

    def _comment_paragraph(self, p, text: str):
        runs = [r for r, _ in _text_runs(p)]
        if runs:
            self._add_comment(_top_level(runs[0], p), _top_level(runs[-1], p), text)

    def _enable_track_changes(self):
        settings = self.doc.settings.element
        if settings.find(qn("w:trackRevisions")) is None:
            el = etree.Element(qn("w:trackRevisions"))
            # Keep schema order loosely: trackRevisions sits early in w:settings.
            anchor = None
            for tag in ("w:zoom", "w:removePersonalInformation", "w:removeDateAndTime",
                        "w:doNotDisplayPageBoundaries", "w:displayBackgroundShape",
                        "w:printPostScriptOverText", "w:printFractionalCharacterWidth",
                        "w:printFormsData", "w:embedTrueTypeFonts", "w:embedSystemFonts",
                        "w:saveSubsetFonts", "w:saveFormsData", "w:mirrorMargins",
                        "w:alignBordersAndEdges", "w:bordersDoNotSurroundHeader",
                        "w:bordersDoNotSurroundFooter", "w:gutterAtTop", "w:hideSpellingErrors",
                        "w:hideGrammaticalErrors", "w:activeWritingStyle", "w:proofState",
                        "w:formsDesign", "w:attachedTemplate", "w:linkStyles",
                        "w:stylePaneFormatFilter", "w:stylePaneSortMethod", "w:documentType",
                        "w:mailMerge", "w:revisionView"):
                found = settings.find(qn(tag))
                if found is not None:
                    anchor = found
            if anchor is not None:
                anchor.addnext(el)
            else:
                settings.insert(0, el)

    def _max_revision_id(self) -> int:
        best = 0
        for el in self.doc.element.body.iter(qn("w:ins"), qn("w:del")):
            try:
                best = max(best, int(el.get(qn("w:id"), "0")))
            except ValueError:
                pass
        return best + 1000  # headroom for revision ids used in headers/footers/footnotes


class _EditError(Exception):
    pass


def _describe(edit: Edit) -> str:
    if edit.kind == "replace":
        return f'replace "{edit.target_text}" with "{edit.new_text}"' if edit.new_text \
            else f'delete "{edit.target_text}"'
    if edit.kind == "insert":
        return f'insert "{edit.new_text}" after "{edit.target_text}"'
    return f'add paragraph "{edit.new_text}"'


# ---------------------------------------------------------------------- run helpers

_TEXT_LEN_TAGS = {qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n", qn("w:noBreakHyphen"): "-"}


def _is_deleted(run, p) -> bool:
    el = run.getparent()
    while el is not None and el is not p:
        if el.tag in (qn("w:del"), qn("w:moveFrom")):
            return True
        el = el.getparent()
    return False


def _run_text(run) -> str:
    parts = []
    for child in run:
        if child.tag == qn("w:t"):
            parts.append(child.text or "")
        elif child.tag in _TEXT_LEN_TAGS:
            parts.append(_TEXT_LEN_TAGS[child.tag])
    return "".join(parts)


def _text_runs(p) -> list[tuple[etree._Element, str]]:
    """Visible runs of a paragraph (skipping text already tracked as deleted)."""
    out = []
    for run in p.iter(qn("w:r")):
        if _is_deleted(run, p):
            continue
        # Runs of a nested paragraph (e.g. text box) belong to that paragraph.
        owner = run.getparent()
        while owner is not None and owner.tag != qn("w:p"):
            owner = owner.getparent()
        if owner is not p:
            continue
        text = _run_text(run)
        if text:
            out.append((run, text))
    return out


def _fold(text: str) -> tuple[str, list[int]]:
    """Normalise text for matching; returns folded text and a map to original indices.

    Runs of whitespace collapse to a single space and are case-folded.
    """
    out_chars: list[str] = []
    index: list[int] = []
    for i, ch in enumerate(text):
        ch = _CHAR_FOLD.get(ch, ch)
        if ch.isspace():
            if out_chars and out_chars[-1] == " ":
                continue
            ch = " "
        out_chars.append(ch.lower())
        index.append(i)
    return "".join(out_chars), index


def _locate(p, target: str) -> tuple[int, int]:
    """Find ``target`` in the paragraph's visible text; returns [start, end) offsets."""
    full = "".join(text for _, text in _text_runs(p))
    pos = full.find(target)
    if pos >= 0:
        return pos, pos + len(target)
    folded, index = _fold(full)
    needle, _ = _fold(target)
    needle = needle.strip()
    if not needle:
        raise _EditError("empty target text")
    pos = folded.find(needle)
    if pos < 0:
        raise _EditError(f'text not found: "{target[:80]}"')
    return index[pos], index[pos + len(needle) - 1] + 1


def _split_run(run, offset: int):
    """Split ``run`` so its first ``offset`` characters stay in it; returns the new right half."""
    right = copy.deepcopy(run)
    left_children = [c for c in run if c.tag != qn("w:rPr")]
    right_children = [c for c in right if c.tag != qn("w:rPr")]
    pos = 0
    for lc, rc in zip(left_children, right_children):
        if lc.tag == qn("w:t"):
            text = lc.text or ""
            length = len(text)
        elif lc.tag in _TEXT_LEN_TAGS:
            text, length = None, 1
        else:
            text, length = None, 0
        if pos + length <= offset:
            # Entirely on the left (zero-length elements before the split stay left).
            right.remove(rc)
        elif pos >= offset:
            run.remove(lc)
        else:  # a w:t that straddles the split point
            cut = offset - pos
            lc.text, rc.text = text[:cut], text[cut:]
            for t in (lc, rc):
                t.set(XML_SPACE, "preserve")
        pos += length
    run.addnext(right)
    return right


def _isolate(p, start: int, end: int) -> list:
    """Split runs so that [start, end) is covered exactly by whole runs; return those runs."""
    pos = 0
    selected = []
    for run, text in _text_runs(p):
        r_start, r_end = pos, pos + len(text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        if r_start < start:
            run = _split_run(run, start - r_start)
            r_start = start
        if r_end > end:
            _split_run(run, end - r_start)
        selected.append(run)
    return selected


def _top_level(el, p):
    """The ancestor of ``el`` that is a direct child of paragraph ``p``."""
    while el.getparent() is not p:
        el = el.getparent()
        if el is None:
            raise _EditError("element not in paragraph")
    return el


def _insert_after(anchor, new_el, p):
    anchor.addnext(new_el)

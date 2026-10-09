"""Generate samples/Sample Target NDA.docx - a deliberately seller-friendly NDA for testing.

    python samples/make_sample_nda.py
"""

from pathlib import Path

import docx
from docx.shared import Pt

PARAGRAPHS = [
    ("title", "CONFIDENTIALITY AGREEMENT"),
    ("", "This Confidentiality Agreement (this “Agreement”) is entered into as of [DATE] "
         "(the “Effective Date”) by and between [TARGET COMPANY] (the “Company”), "
         "acting through its financial advisor, Harbor Point Capital Partners, and Unifi Aviation, LLC "
         "(“Recipient”)."),
    ("", "1. Confidential Information. “Confidential Information” means all information, "
         "including trade secrets, concerning the Company furnished to Recipient before, on or after the "
         "date of this Agreement, together with all analyses, compilations and notes prepared by Recipient."),
    ("", "2. Exceptions. Confidential Information does not include information that (i) is or becomes "
         "generally available to the public other than as a result of disclosure by Recipient, or (ii) "
         "was within Recipient’s possession prior to its being furnished, as evidenced by written "
         "documentation."),
    ("", "3. Representatives. “Representatives” means Recipient’s officers, directors, "
         "employees and attorneys to whom Confidential Information is provided by you."),
    ("", "4. Use and Disclosure. Recipient shall hold the Confidential Information in strict confidence and "
         "shall disclose it only to Representatives who need to know such information for the purpose of "
         "evaluating a possible transaction with the Company (the “Transaction”). Recipient shall "
         "be responsible for any breach of this Agreement by any person to whom it discloses Confidential "
         "Information."),
    ("", "5. Compelled Disclosure. If Recipient is requested or required by law to disclose any Confidential "
         "Information, Recipient shall provide the Company with immediate prior written notice so that the "
         "Company may seek a protective order, and Recipient shall disclose only that portion which it is "
         "advised by counsel is legally required."),
    ("", "6. Return of Materials. Upon the Company’s demand, or at the conclusion of the parties’ "
         "discussions, Recipient shall return all Confidential Information, including all copies, written "
         "notes, photographs, and memoranda, and an officer of Recipient shall certify such return in writing."),
    ("", "7. Non-Solicitation. For a period of three (3) years from the Effective Date, neither Recipient nor "
         "its Representatives shall, directly or indirectly, solicit for employment or hire any current or "
         "former employee of the Company, or contact any customer or supplier of the Company."),
    ("", "8. Remedies. Recipient acknowledges that money damages would not be a sufficient remedy for any "
         "breach of this Agreement and that the Company shall be entitled to specific performance and "
         "injunctive relief. Recipient shall indemnify and hold harmless the Company from any losses "
         "arising out of any breach by Recipient or its Representatives."),
    ("", "9. Governing Law. This Agreement shall be governed by the laws of the State of California, and "
         "each party submits to the exclusive jurisdiction of the courts located in San Francisco County."),
    ("", "10. Term. The obligations under this Agreement shall survive indefinitely."),
    ("", "11. Definitive Agreement. The parties agree to negotiate in good faith to enter into a definitive "
         "agreement with respect to the Transaction."),
    ("", "IN WITNESS WHEREOF, the parties have executed this Agreement as of the Effective Date."),
]


def build(path: Path) -> None:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for kind, text in PARAGRAPHS:
        p = document.add_paragraph()
        if kind == "title":
            p.add_run(text).bold = True
            continue
        # Split the clause heading into its own bold run, like real NDAs.
        head, _, rest = text.partition(". ") if text[:1].isdigit() else ("", "", text)
        if head:
            heading, _, body_start = rest.partition(". ")
            p.add_run(f"{head}. ")
            p.add_run(f"{heading}.").bold = True
            rest = f" {body_start}"
        p.add_run(rest)
    document.save(str(path))


if __name__ == "__main__":
    out = Path(__file__).with_name("Sample Target NDA.docx")
    build(out)
    print(f"wrote {out}")
